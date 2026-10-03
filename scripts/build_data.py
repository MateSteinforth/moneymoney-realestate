#!/usr/bin/env python3
"""Build the index data consumed by RealEstateTracker.lua.

Downloads free, official/open sources and writes two JSON files:

  data/indices.json  quarterly price-index series per (area, property type)
  data/de_plz.json   German postcode (PLZ) -> area id

Sources
  GREIX            Kiel Institut / ECONtribute, transaction-based city indices (greix.de)
  Destatis         Häuserpreisindex nach siedlungsstrukturellen Kreistypen (Statistischer Bericht 61262)
  BBSR             Raumgliederungen-Referenzen (district -> settlement type)
  OpenPLZ / OSM    streets.csv (postcode -> municipality key), ODbL
  Eurostat         prc_hpi_q, house price index per country

Requires: python3, openpyxl
"""

import csv
import io
import json
import math
import os
import re
import sys
import urllib.request
from collections import Counter, defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
UA = {"User-Agent": "Mozilla/5.0 (moneymoney-realestate data builder)"}
START = 1975 * 4  # drop anything older than 1975-Q1 to keep the file small

GREIX_CITIES_XLSX = "https://api.greixx.net/api-v1/cities/metrics/download/xlsx/all/?"
DESTATIS_PAGE = "https://www.destatis.de/DE/Themen/Wirtschaft/Preise/Baupreise-Immobilienpreisindex/_inhalt.html"
DESTATIS_XLSX_FALLBACK = (
    "https://www.destatis.de/DE/Themen/Wirtschaft/Preise/Baupreise-Immobilienpreisindex/Publikationen/"
    "Downloads-Bau-und-Immobilienpreisindex/statistischer-bericht-wohnimmobilien-kreistypen-revisionsdifferenzen-5612603.xlsx"
    "?__blob=publicationFile"
)
BBSR_PAGE = "https://www.bbsr.bund.de/BBSR/DE/forschung/raumbeobachtung/Raumabgrenzungen/downloads/download-referenzen.html"
OPENPLZ_STREETS = "https://raw.githubusercontent.com/openpotato/openplzapi.data/main/src/de/osm/streets.csv"
EUROSTAT_HPI = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/prc_hpi_q?purchase=TOTAL&unit=I15_Q"

TYPES = ("etw", "efh", "mfh")
GREIX_TYPES = {"Apartment": "etw", "Single-family house": "efh", "Multi-family house": "mfh"}
# Destatis publishes no multi-family index; apartments are the closest proxy.
DESTATIS_TYPE = {"etw": "etw", "efh": "efh", "mfh": "etw"}

# GREIX city -> official municipality key (AGS, 8 digits) or district prefix (5 digits).
GREIX_AREAS = {
    "Berlin": "11000000",
    "Hamburg": "02000000",
    "Munich": "09162000",
    "Cologne": "05315000",
    "Frankfurt": "06412000",
    "Duesseldorf": "05111000",
    "Stuttgart": "08111000",
    "Dortmund": "05913000",
    "Duisburg": "05112000",
    "Bochum": "05911000",
    "Bonn": "05314000",
    "Muenster": "05515000",
    "Hamm": "05915000",
    "Bocholt": "05554008",
    "Kreis Mettmann": "05158",
    "Rhein-Erft-Kreis": "05362",
    "Chemnitz": "14511000",
    "Dresden": "14612000",
    "Leipzig": "14713000",
    "Erfurt": "16051000",
    "Karlsruhe": "08212000",
    "Luebeck": "01003000",
    "Potsdam": "12054000",
    "Wiesbaden": "06414000",
    "Landkreis Ludwigslust-Parchim": "13076",
}
GREIX_LABELS = {"Munich": "München", "Cologne": "Köln", "Frankfurt": "Frankfurt am Main",
                "Duesseldorf": "Düsseldorf", "Muenster": "Münster", "Luebeck": "Lübeck"}

# The seven cities Destatis reports as a separate "Metropolen" type.
METRO_DISTRICTS = {"11000", "02000", "09162", "05315", "06412", "08111", "05111"}
KREISTYP_LABELS = {
    "de-metro": "Metropolen (Top 7)",
    "de-kt1": "Kreisfreie Großstädte",
    "de-kt2": "Städtische Kreise",
    "de-kt3": "Ländliche Kreise mit Verdichtungsansätzen",
    "de-kt4": "Dünn besiedelte ländliche Kreise",
}

EUROSTAT_SKIP = re.compile(r"^(EU|EA)")


def log(*a):
    print(*a, file=sys.stderr)


def fetch(url, timeout=300):
    log("GET", url)
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r:
        return r.read()


def open_xlsx(blob):
    return openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)


def q_index(year, quarter):
    return int(year) * 4 + int(quarter) - 1


def q_label(qi):
    return f"{qi // 4}-Q{qi % 4 + 1}"


def slug(s):
    s = s.lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


# --------------------------------------------------------------------------- series helpers


def interpolate(points):
    """Fill gaps between known quarters by log-linear interpolation."""
    if not points:
        return {}
    qs = sorted(points)
    out = {}
    for a, b in zip(qs, qs[1:]):
        va, vb = points[a], points[b]
        for q in range(a, b):
            t = (q - a) / (b - a)
            out[q] = math.exp(math.log(va) * (1 - t) + math.log(vb) * t)
    out[qs[-1]] = points[qs[-1]]
    return out


def chain(primary, fallback):
    """Extend `primary` backwards and forwards using the growth rates of `fallback`."""
    if not primary:
        return dict(fallback or {})
    out = dict(primary)
    if not fallback:
        return out
    first, last = min(primary), max(primary)
    if first in fallback:
        for q in fallback:
            if q < first:
                out[q] = primary[first] * fallback[q] / fallback[first]
    if last in fallback:
        for q in fallback:
            if q > last:
                out[q] = primary[last] * fallback[q] / fallback[last]
    return out


def span(series):
    return (q_label(min(series)), q_label(max(series))) if series else None


# --------------------------------------------------------------------------- sources


def load_greix():
    wb = open_xlsx(fetch(GREIX_CITIES_XLSX))
    rows = wb.worksheets[0].iter_rows(values_only=True)
    header = next(rows)
    col = {name: i for i, name in enumerate(header)}
    quarterly, annual = defaultdict(dict), defaultdict(dict)
    for r in rows:
        if r[col["Inflation"]] != 0 or r[col["Index"]] is None:
            continue  # keep nominal values only
        ptype = GREIX_TYPES.get(r[col["Property_Type"]])
        if not ptype:
            continue
        key = (r[col["City"]], ptype)
        year, quarter, value = r[col["Year"]], r[col["Quarter"]], float(r[col["Index"]])
        if quarter:
            quarterly[key][q_index(year, quarter)] = value
        else:
            annual[key][q_index(year, 2)] = value  # place annual values mid-year

    series = {}
    for key in set(quarterly) | set(annual):
        q = interpolate(quarterly.get(key, {}))
        a = interpolate(annual.get(key, {}))
        # Quarterly and annual series are estimated separately; quarterly wins,
        # annual growth only extends it where no quarterly data exists.
        series[key] = chain(q, a)
    log(f"GREIX: {len(series)} city/type series")
    return series


def find_link(page_url, pattern):
    html = fetch(page_url).decode("utf-8", "ignore")
    links = re.findall(r'href="([^"]*' + pattern + r'[^"]*)"', html)
    if not links:
        return None
    link = sorted(set(links))[-1].replace("&amp;", "&")
    return link if link.startswith("http") else "https://" + page_url.split("/")[2] + link


def load_destatis():
    url = find_link(DESTATIS_PAGE, r"statistischer-bericht-wohnimmobilien-kreistypen[^\"]*\.xlsx") or DESTATIS_XLSX_FALLBACK
    wb = open_xlsx(fetch(url))
    areas = ["de-metro", "de-kt1", "de-kt2", "de-kt3", "de-kt4"]
    series = {}
    for sheet, ptype in (("61262-b01", "etw"), ("61262-b02", "efh")):
        data = {a: {} for a in areas}
        for row in wb[sheet].iter_rows(values_only=True):
            m = re.match(r"(\d)\. Vierteljahr (\d{4})", str(row[0] or ""))
            if not m:
                continue
            qi = q_index(m.group(2), m.group(1))
            for area, v in zip(areas, row[1:6]):
                if isinstance(v, (int, float)):
                    data[area][qi] = float(v)
        for area in areas:
            series[(area, ptype)] = data[area]
    log(f"Destatis: {len(series)} series, {span(series[('de-metro', 'etw')])}")
    return series


def load_bbsr():
    url = find_link(BBSR_PAGE, r"raumgliederungen-referenzen-\d{4}\.xlsx")
    if not url:
        raise SystemExit("BBSR reference file not found")
    wb = open_xlsx(fetch(url))
    rows = wb["Kreisreferenz"].iter_rows(values_only=True)
    header = next(rows)
    next(rows)  # long column descriptions
    krs = next(i for i, h in enumerate(header) if re.fullmatch(r"KRS\d{4}", str(h)))
    ktu = next(i for i, h in enumerate(header) if re.fullmatch(r"KTU\d{4}", str(h)))
    out = {}
    for r in rows:
        if r[krs] is None or r[ktu] is None:
            continue
        district = str(r[krs]).zfill(8)[:5]
        out[district] = "de-metro" if district in METRO_DISTRICTS else f"de-kt{int(r[ktu])}"
    log(f"BBSR: {len(out)} districts")
    return out


def load_plz_to_ags():
    text = fetch(OPENPLZ_STREETS).decode("utf-8")
    counts = defaultdict(Counter)
    for row in csv.DictReader(io.StringIO(text)):
        plz, ags = row["PostalCode"], row["RegionalKey"]
        if re.fullmatch(r"\d{5}", plz or "") and re.fullmatch(r"\d{8}", ags or ""):
            counts[plz][ags] += 1
    out = {plz: c.most_common(1)[0][0] for plz, c in counts.items()}
    log(f"OpenPLZ: {len(out)} postcodes")
    return out


def load_eurostat():
    d = json.loads(fetch(EUROSTAT_HPI))
    dims = d["id"]
    sizes = d["size"]
    cats = {k: d["dimension"][k]["category"] for k in dims}
    geo_index = cats["geo"]["index"]
    time_index = cats["time"]["index"]
    inv_geo = {v: k for k, v in geo_index.items()}
    inv_time = {v: k for k, v in time_index.items()}
    stride_geo = math.prod(sizes[dims.index("geo") + 1:])
    stride_time = math.prod(sizes[dims.index("time") + 1:])
    series = defaultdict(dict)
    for flat, value in d["value"].items():
        flat = int(flat)
        geo = inv_geo[(flat // stride_geo) % sizes[dims.index("geo")]]
        period = inv_time[(flat // stride_time) % sizes[dims.index("time")]]
        m = re.fullmatch(r"(\d{4})-Q(\d)", period)
        if m and not EUROSTAT_SKIP.match(geo):
            series[geo][q_index(*m.groups())] = float(value)
    labels = cats["geo"].get("label", {})
    log(f"Eurostat: {len(series)} countries")
    return dict(series), labels


# --------------------------------------------------------------------------- assembly


def build():
    greix = load_greix()
    destatis = load_destatis()
    districts = load_bbsr()
    plz_ags = load_plz_to_ags()
    eurostat, eu_labels = load_eurostat()

    series, areas = {}, {}

    def put(area, ptype, values, sources):
        values = {q: v for q, v in values.items() if q >= START}
        if not values:
            return
        start, end = min(values), max(values)
        series[f"{area}:{ptype}"] = {
            "start": q_label(start),
            "end": q_label(end),
            "v": [round(values[q], 3) for q in range(start, end + 1)],
            "src": sources,
        }

    # National fallback: GREIX composite (long history across all GREIX cities).
    national = {}
    for t in TYPES:
        national[t] = greix.get(("GREIX", t), {})
        put("de", t, national[t], [f"GREIX Gesamtindex {span(national[t])[0]}–{span(national[t])[1]}"])
    areas["de"] = {"label": "Deutschland (GREIX Gesamtindex)"}

    # Destatis settlement types, back-cast with the national GREIX composite.
    destatis_chained = {}
    for area, label in KREISTYP_LABELS.items():
        areas[area] = {"label": f"Destatis {label}"}
        for t in TYPES:
            primary = destatis[(area, DESTATIS_TYPE[t])]
            full = chain(primary, national[t])
            destatis_chained[(area, t)] = full
            src = [f"Destatis Häuserpreisindex {label} ({'ETW' if DESTATIS_TYPE[t] == 'etw' else 'EFH'}) {span(primary)[0]}–{span(primary)[1]}",
                   "davor: GREIX Gesamtindex"]
            put(area, t, full, src)

    # GREIX cities, extended with the Destatis type of their district.
    greix_by_key = {}
    for city, key in GREIX_AREAS.items():
        district = key[:5]
        kt = districts.get(district)
        if not kt:
            log(f"WARN: no BBSR district for {city} ({key})")
            continue
        area = "greix-" + slug(city)
        areas[area] = {"label": f"GREIX {GREIX_LABELS.get(city, city)}"}
        greix_by_key[key] = area
        for t in TYPES:
            primary = greix.get((city, t), {})
            fallback = destatis_chained[(kt, t)]
            src = []
            if primary:
                src.append(f"GREIX {GREIX_LABELS.get(city, city)} {t.upper()} {span(primary)[0]}–{span(primary)[1]}")
                src.append(f"fortgeschrieben/zurückgerechnet mit Destatis {KREISTYP_LABELS[kt]}")
            else:
                src.append(f"GREIX hat keinen {t.upper()}-Index für {GREIX_LABELS.get(city, city)}: Destatis {KREISTYP_LABELS[kt]}")
            put(area, t, chain(primary, fallback), src)

    # Eurostat country indices (all dwellings) for objects outside Germany.
    for geo, values in sorted(eurostat.items()):
        area = "eu-" + geo
        areas[area] = {"label": f"Eurostat HPI {eu_labels.get(geo, geo)}"}
        put(area, "all", values, [f"Eurostat prc_hpi_q {geo} {span(values)[0]}–{span(values)[1]}"])

    # PLZ -> area. GREIX match on municipality first, then district; else Destatis type.
    plz_area = {}
    for plz, ags in sorted(plz_ags.items()):
        area = greix_by_key.get(ags) or greix_by_key.get(ags[:5]) or districts.get(ags[:5])
        if area:
            plz_area[plz] = area
        else:
            log(f"WARN: PLZ {plz} -> {ags} has no area")
    # 3-digit prefix fallback for postcodes that are missing from the street data.
    prefix = defaultdict(Counter)
    for plz, area in plz_area.items():
        prefix[plz[:3]][area] += 1
    plz_prefix = {p: c.most_common(1)[0][0] for p, c in sorted(prefix.items())}

    # No timestamps: the files only change (and get committed) when the data changes.
    indices = {
        "version": 1,
        "data_through": max(s["end"] for s in series.values()),
        "attribution": [
            "GREIX – German Real Estate Index, Kiel Institut für Weltwirtschaft / ECONtribute, greix.de",
            "Statistisches Bundesamt (Destatis), Häuserpreisindex nach Kreistypen, dl-de/by-2-0",
            "BBSR, Raumgliederungen (Siedlungsstrukturelle Kreistypen), dl-de/by-2-0",
            "Eurostat, House price index (prc_hpi_q), CC BY 4.0",
        ],
        "areas": areas,
        "series": series,
    }
    plz = {
        "version": 1,
        "attribution": ["© OpenStreetMap contributors (ODbL) via OpenPLZ API data", "BBSR Raumgliederungen, dl-de/by-2-0"],
        "plz": plz_area,
        "prefix": plz_prefix,
    }

    os.makedirs(DATA, exist_ok=True)
    for name, obj in (("indices.json", indices), ("de_plz.json", plz)):
        path = os.path.join(DATA, name)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(obj, f, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
            f.write("\n")
        log(f"wrote {path} ({os.path.getsize(path) // 1024} KB)")


if __name__ == "__main__":
    build()
